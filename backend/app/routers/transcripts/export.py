"""transcripts.py 가 792줄로 커져 부분 분리.  이 파일의 모든 route 는
transcripts._core.router (prefix='/api/transcripts') 에 직접 등록.  main.py
의 include_router 는 transcripts 패키지의 단일 router 만 부르므로 분할은
internal-only."""
from __future__ import annotations

import asyncio
import logging
import os
import tempfile

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from ... import models, schemas
from ...auth import get_current_user, require_admin
from ...config import settings
from ...database import get_db
from ...transcribe.runner import run_transcription

from ._core import router, _resolve_orphan_session

@router.get("/{transcript_id}/export.hwpx")
async def export_transcript_hwpx(
    transcript_id: str,
    include: str = "summary",
    message_ids: str = "",
    mask_pii: bool = False,
    db: AsyncSession = Depends(get_db),
    user: models.User = Depends(get_current_user),
):
    """회의록을 HWPX (한컴 오픈 XML) 로 다운로드. include / message_ids /
    mask_pii 동작은 export.docx 와 동일."""
    import urllib.parse
    from fastapi.responses import Response
    from ... import hwpx_export, pii_mask

    orphan_sess = await _resolve_orphan_session(db, transcript_id, user.id)
    if orphan_sess is not None:
        sess = orphan_sess
        title = sess.title or "회의록"
    else:
        tr = (
            await db.execute(
                select(models.Transcript).where(
                    models.Transcript.id == transcript_id,
                    models.Transcript.user_id == user.id,
                )
            )
        ).scalar_one_or_none()
        if tr is None:
            raise HTTPException(404, "transcript not found")
        if not tr.session_id:
            raise HTTPException(409, "전사가 아직 완료되지 않았습니다")
        sess = await db.scalar(
            select(models.Session).where(
                models.Session.id == tr.session_id,
                models.Session.user_id == user.id,
            )
        )
        if sess is None:
            raise HTTPException(404, "연결된 채팅 세션을 찾을 수 없습니다")
        title = tr.source_filename or sess.title or "회의록"

    msg_rows = (
        await db.execute(
            select(models.Message)
            .where(models.Message.session_id == sess.id)
            .order_by(models.Message.created_at.asc())
        )
    ).scalars().all()

    picked_ids: set[str] | None = None
    if message_ids.strip():
        picked_ids = {x.strip() for x in message_ids.split(",") if x.strip()}
    if picked_ids is not None:
        selected = [m for m in msg_rows if m.id in picked_ids]
    elif include == "all":
        selected = list(msg_rows)
    else:
        selected = [m for m in msg_rows if m.role == "assistant"]

    paragraphs = ["회의록", title, ""]
    for m in selected:
        text = (m.content or "").strip()
        if not text:
            continue
        if mask_pii:
            text = pii_mask.mask(text)
        paragraphs.append("[요약]" if m.role == "assistant" else "[원문/메모]")
        paragraphs.append(text)
        paragraphs.append("")

    data = hwpx_export.build_hwpx(title, paragraphs)
    safe = urllib.parse.quote((title or "회의록").replace("/", "_"))
    return Response(
        content=data,
        media_type="application/hwp+zip",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{safe}.hwpx"
        },
    )


@router.get("/{transcript_id}/export.docx")
async def export_transcript_docx(
    transcript_id: str,
    include: str = "summary",
    message_ids: str = "",
    mask_pii: bool = False,
    db: AsyncSession = Depends(get_db),
    user: models.User = Depends(get_current_user),
):
    """Render the transcript's linked chat session as a Korean
    회의록 DOCX and stream it back.

    Query params control which messages land in the document:
      · `include=summary` (default) — assistant messages only, raw
        transcript hidden. Matches the common "회의록 = 요약만" need.
      · `include=all` — every message including the raw transcript.
      · `message_ids=<id>,<id>,...` — explicit pick from the
        selection modal in the UI. Overrides `include` when set.
    """
    import io
    import urllib.parse

    # Orphan rows export from the linked Session directly (no
    # Transcript row to look up). `tr` stays None for that path so the
    # title-page rendering below has to fall back to session attrs.
    tr: models.Transcript | None = None
    orphan_sess = await _resolve_orphan_session(db, transcript_id, user.id)
    if orphan_sess is not None:
        sess = orphan_sess
    else:
        tr = (
            await db.execute(
                select(models.Transcript).where(
                    models.Transcript.id == transcript_id,
                    models.Transcript.user_id == user.id,
                )
            )
        ).scalar_one_or_none()
        if tr is None:
            raise HTTPException(404, "transcript not found")
        if not tr.session_id:
            raise HTTPException(409, "전사가 아직 완료되지 않았습니다")
        sess = await db.scalar(
            select(models.Session).where(
                models.Session.id == tr.session_id,
                models.Session.user_id == user.id,
            )
        )
        if sess is None:
            raise HTTPException(404, "연결된 채팅 세션을 찾을 수 없습니다")
    msg_rows = (
        await db.execute(
            select(models.Message)
            .where(models.Message.session_id == sess.id)
            .order_by(models.Message.created_at.asc())
        )
    ).scalars().all()

    # Decide which messages to include. Explicit IDs win; otherwise
    # the `include` mode picks a sensible default.
    picked_ids: set[str] | None = None
    if message_ids.strip():
        picked_ids = {
            x.strip() for x in message_ids.split(",") if x.strip()
        }
    selected_msgs: list[models.Message] = []
    if picked_ids is not None:
        selected_msgs = [m for m in msg_rows if m.id in picked_ids]
    elif include == "all":
        selected_msgs = list(msg_rows)
    else:  # "summary" (default)
        selected_msgs = [m for m in msg_rows if m.role == "assistant"]

    try:
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError as exc:
        raise HTTPException(
            500,
            "python-docx 가 설치돼 있지 않습니다. "
            "백엔드 의존성 설치를 확인하세요.",
        ) from exc

    doc = Document()

    # Title page
    title = doc.add_heading("회의록", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_text = (
        (tr.source_filename if tr else None) or sess.title or "녹음"
    )
    run = subtitle.add_run(title_text)
    run.bold = True
    run.font.size = Pt(14)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_bits: list[str] = []
    created_at = (tr.created_at if tr else None) or sess.created_at
    if created_at:
        meta_bits.append(f"작성 일시: {created_at.strftime('%Y-%m-%d %H:%M')}")
    if tr and tr.duration_sec:
        m_, s_ = divmod(int(tr.duration_sec), 60)
        h_, m_ = divmod(m_, 60)
        if h_:
            meta_bits.append(f"녹음 길이: {h_}h {m_:02d}m {s_:02d}s")
        else:
            meta_bits.append(f"녹음 길이: {m_}m {s_:02d}s")
    if tr and tr.language:
        meta_bits.append(f"언어: {tr.language}")
    if tr and tr.diarized:
        meta_bits.append("화자 분리: 적용")
    if tr is None:
        meta_bits.append("보관됨 (음원 삭제)")
    meta_run = meta.add_run("  ·  ".join(meta_bits))
    meta_run.italic = True
    meta_run.font.size = Pt(10)

    doc.add_paragraph()  # spacer

    if not selected_msgs:
        doc.add_paragraph(
            "선택된 본문이 없습니다. 채팅창에서 회의록에 담을 메시지를 "
            "선택한 뒤 다시 받아주세요."
        )
    else:
        # Each picked message renders as one section. Numbering only
        # appears when there's more than one of the same role so a
        # single-summary export reads cleanly as just "요약".
        role_counts = {"assistant": 0, "user": 0}
        for m in selected_msgs:
            role_counts[m.role] += 1
        a_idx = 0
        u_idx = 0
        from ... import pii_mask
        for m in selected_msgs:
            content = (m.content or "").strip()
            if not content:
                continue
            if mask_pii:
                content = pii_mask.mask(content)
            if m.role == "assistant":
                a_idx += 1
                label = "요약"
                if role_counts["assistant"] > 1:
                    label = f"요약 {a_idx}"
                doc.add_heading(label, level=1)
            else:
                u_idx += 1
                # User-side: distinguish the original transcript (first
                # one) from later notes the user typed into chat.
                if u_idx == 1 and role_counts["user"] >= 1:
                    label = "전체 전사"
                else:
                    label = f"메모 {u_idx - 1}"
                doc.add_heading(label, level=1)
            for para in content.split("\n\n"):
                line = para.strip()
                if not line:
                    continue
                if line.startswith("# "):
                    doc.add_heading(line[2:].strip(), level=2)
                elif line.startswith("## "):
                    doc.add_heading(line[3:].strip(), level=2)
                elif line.startswith("### "):
                    doc.add_heading(line[4:].strip(), level=3)
                else:
                    p = doc.add_paragraph()
                    p.add_run(line)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    base = (tr.source_filename or sess.title or "회의록").rsplit(".", 1)[0]
    leaf = f"{base} 회의록.docx"
    ascii_fallback = (
        leaf.encode("ascii", errors="replace")
        .decode("ascii")
        .replace('"', "_")
    )
    encoded = urllib.parse.quote(leaf, safe="")
    disposition = (
        f'attachment; filename="{ascii_fallback}"; '
        f"filename*=UTF-8''{encoded}"
    )
    from fastapi.responses import StreamingResponse

    return StreamingResponse(
        buf,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        headers={"Content-Disposition": disposition},
    )


